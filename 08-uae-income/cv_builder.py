#!/usr/bin/env python3
"""Shared ATS-safe CV builder. Produces .docx and .pdf from one content dict.

Layout lives here; wording lives in the per-variant build scripts, so a new
variant is a content file rather than a copy of the layout code.

ATS rules, from 05-career-context/cv-template-ats.md:
  single column, no tables, contact details in the BODY not the header,
  standard headings, standard bullets, Month Year - Month Year dates.

The PDF is laid out directly rather than converted from the .docx, because
LibreOffice is non-functional in this environment (it fails to load even a
plain text file). Both outputs carry real text layers; neither is exported
from a design tool, which is the case application-tailor warns about.

Contact details are never hard-coded here or in a variant script - they come
from CV_EMAIL and CV_PHONE. application-history.md: "Contact details are not
stored here." Committing them would put them in git history permanently.

Deps: pip install python-docx reportlab

Content dict shape:
    name, tagline, location, links, summary, skills, tools  -> str
    roles     -> [{title, employer, dates, note|None, bullets: [str]}]
    education -> [(bold_title, [detail lines])]
    languages -> [str]
"""
import os

EMAIL = os.environ.get("CV_EMAIL", "[EMAIL]")
PHONE = os.environ.get("CV_PHONE", "[PHONE]")


def contact_line():
    return f"{EMAIL}  |  {PHONE}"


def placeholders_remain():
    return EMAIL.startswith("[") or PHONE.startswith("[")


# ------------------------------------------------------------------ DOCX
def build_docx(c, path):
    from docx import Document
    from docx.shared import Pt, Inches

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05

    def para(text="", size=10.5, bold=False, italic=False, after=4, before=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        if text:
            r = p.add_run(text)
            r.font.size = Pt(size)
            r.bold = bold
            r.italic = italic
        return p

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(11)

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Inches(0.25)
        p.add_run(text).font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(c["name"])
    r.bold = True
    r.font.size = Pt(19)

    para(c["tagline"], size=11.5, after=6)
    para(c["location"], size=10, after=1)
    para(contact_line(), size=10, after=1)
    para(c["links"], size=10, after=2)

    heading("Summary")
    para(c["summary"])

    heading("Experience")
    for role in c["roles"]:
        pp = doc.add_paragraph()
        pp.paragraph_format.space_before = Pt(9)
        pp.paragraph_format.space_after = Pt(0)
        rr = pp.add_run(role["title"])
        rr.bold = True
        rr.font.size = Pt(11)

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        p2.add_run(role["employer"]).font.size = Pt(10.5)

        p3 = doc.add_paragraph()
        p3.paragraph_format.space_after = Pt(4)
        r3 = p3.add_run(role["dates"] + (f"  |  {role['note']}" if role.get("note") else ""))
        r3.font.size = Pt(10)
        r3.italic = True

        for b in role["bullets"]:
            bullet(b)

    heading("Education")
    for title, lines in c["education"]:
        para(title, bold=True, after=1)
        for i, ln in enumerate(lines):
            para(ln, size=10, after=6 if i == len(lines) - 1 else 1)

    heading("Skills")
    para(c["skills"])

    heading("Languages")
    for i, ln in enumerate(c["languages"]):
        para(ln, after=1 if i < len(c["languages"]) - 1 else 4)

    heading("Tools")
    para(c["tools"])

    doc.save(path)
    return path


# ------------------------------------------------------------------- PDF
def build_pdf(c, path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                    ListFlowable, ListItem)

    # Tuned to land on two pages. cv-template-ats.md: "Two pages is fine at
    # 22 years. Do not compress to one and lose the evidence." Three pages
    # with a short orphan is worse than either.
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=9.5, leading=11.8, spaceAfter=3)
    small = ParagraphStyle("small", parent=body, fontSize=9, leading=11, spaceAfter=1)
    ital = ParagraphStyle("ital", parent=small, fontName="Helvetica-Oblique", spaceAfter=3)
    name = ParagraphStyle("name", parent=body, fontName="Helvetica-Bold", fontSize=18,
                          leading=20, spaceAfter=2)
    tag = ParagraphStyle("tag", parent=body, fontSize=10.5, leading=12.5, spaceAfter=5)
    head = ParagraphStyle("head", parent=body, fontName="Helvetica-Bold", fontSize=10,
                          leading=12, spaceBefore=9, spaceAfter=3)
    rtitle = ParagraphStyle("rtitle", parent=body, fontName="Helvetica-Bold", fontSize=10,
                            leading=12, spaceBefore=6.5, spaceAfter=1)
    bul = ParagraphStyle("bul", parent=body, fontSize=9.5, leading=11.8, spaceAfter=2.5)

    doc = SimpleDocTemplate(path, pagesize=A4,
                            leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                            topMargin=0.6 * inch, bottomMargin=0.6 * inch,
                            title=f"{c['name']} - CV", author=c["name"])

    def esc(t):
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    f = [Paragraph(esc(c["name"]), name),
         Paragraph(esc(c["tagline"]), tag),
         Paragraph(esc(c["location"]), small),
         Paragraph(esc(contact_line()), small),
         Paragraph(esc(c["links"]), small),
         Paragraph("SUMMARY", head), Paragraph(esc(c["summary"]), body),
         Paragraph("EXPERIENCE", head)]

    for role in c["roles"]:
        f.append(Paragraph(esc(role["title"]), rtitle))
        f.append(Paragraph(esc(role["employer"]), small))
        f.append(Paragraph(esc(role["dates"] +
                               (f"  |  {role['note']}" if role.get("note") else "")), ital))
        f.append(ListFlowable(
            [ListItem(Paragraph(esc(b), bul), leftIndent=14, value="bulletchar")
             for b in role["bullets"]],
            bulletType="bullet", bulletFontSize=7, leftIndent=14, start="•"))

    f.append(Paragraph("EDUCATION", head))
    for title, lines in c["education"]:
        f.append(Paragraph(f"<b>{esc(title)}</b>", body))
        for ln in lines:
            f.append(Paragraph(esc(ln), small))
        f.append(Spacer(1, 5))

    f += [Paragraph("SKILLS", head), Paragraph(esc(c["skills"]), body),
          Paragraph("LANGUAGES", head)]
    for ln in c["languages"]:
        f.append(Paragraph(esc(ln), small))
    f += [Paragraph("TOOLS", head), Paragraph(esc(c["tools"]), body)]

    doc.build(f)
    return path


def build_both(content, basename):
    print("Written:", build_docx(content, f"{basename}.docx"))
    print("Written:", build_pdf(content, f"{basename}.pdf"))
    if placeholders_remain():
        print("\n!! Contact details are placeholders. Re-run with:")
        print(f'   CV_EMAIL="..." CV_PHONE="..." python3 <script>')
