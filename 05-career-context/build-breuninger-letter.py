#!/usr/bin/env python3
"""Breuninger cover letter. Plain, single column, no tables."""

from docx import Document
from docx.shared import Pt, Inches

EMAIL = "hello@sarahradwan.me"
PHONE = "+971 52 870 2003"
OUT = ("/tmp/claude-0/-home-user-sue-eu/a614d373-7f21-52dd-b6d2-bd7c13427802/scratchpad/"
       "SaraRadwan_Anschreiben_Breuninger.docx")

doc = Document()
for s in doc.sections:
    s.top_margin = Inches(0.8); s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.9); s.right_margin = Inches(0.9)
st = doc.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(8)
st.paragraph_format.line_spacing = 1.12


def para(text, size=10.5, bold=False, italic=False, after=8):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    p.paragraph_format.space_after = Pt(after)
    return p


p = doc.add_paragraph()
r = p.add_run("Sara Radwan"); r.bold = True; r.font.size = Pt(14)
p.paragraph_format.space_after = Pt(1)
para(f"{EMAIL} | {PHONE} | sarahradwan.me/case-studies", size=9.5, after=14)

para("Senior Art Director Image (m/w/d), Stuttgart", bold=True, after=10)

para("Hallo Sabrina,", after=8)

para("diese ersten Zeilen schreibe ich auf Deutsch, damit du mein Niveau direkt siehst: ich bin "
     "auf A2 und lerne weiter. Ab hier auf Englisch.")

para("What made me read the whole posting was the team description. An Ideation Team where AI sits "
     "as a named discipline beside Grafik and Art Direction, rather than as an experiment someone "
     "runs on the side. That is already how I work. I use generative image tools in live production, "
     "build prompting workflows, and write AI usage guidelines into client brand systems, so the "
     "question of where AI belongs is settled before a project starts instead of argued about "
     "halfway through it.")

para("The rest of the role is where most of my career has been. Fifteen of my twenty-two years are "
     "editorial. At CPI Media Group in Dubai I held image and design direction across more than 20 "
     "consumer and B2B titles in fashion, beauty and lifestyle, and I was also Fashion and Beauty "
     "Editor, so the content was mine as well as the look. I directed cover and interview shoots: "
     "briefing photographers, setting the visual direction, making the calls on set. A cover is the "
     "most exposed image a title produces.")

para("I should be precise about the limit of that. My shoot experience is editorial, covers and "
     "portraits, not fashion campaign productions with stylists and set designers at your scale. I "
     "have steered production partners for years and I would run those, but I would rather you knew "
     "the shape of my experience than assumed it.")

para("Print production I know very well. A national daily newspaper redesign at Al Arab, where the "
     "deadline arrives every day. A 100-title curriculum series for the UAE Ministry of Education "
     "across five subjects and three scripts, where I built the typographic system and the templates "
     "that let a large team hold it. That part of the craft is quietly disappearing and I still have "
     "all of it.")

para("Two practical things. My German is A2 and improving, and I read it better than I speak it. If "
     "the role needs confident German from day one, tell me and I will understand. And I am not an "
     "EU citizen, so I would need a work permit. Germany does not require an employer to be "
     "pre-registered for that, and I would carry as much of the process as I can.")

para("Sara Radwan", after=1)
para("sarahradwan.me", size=9.5, after=2)

doc.save(OUT)

from docx import Document as D
d = D(OUT)
assert len(d.tables) == 0
text = "\n".join(p.text for p in d.paragraphs)
assert "—" not in text, "em dash found"
print(f"Saved: {OUT}")
print(f"Tables: {len(d.tables)} | words: {len(text.split())} | em dashes: 0")
